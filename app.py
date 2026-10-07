"""
Leuchten GIS Dashboard v4.2
Lichtplanungsbuero - Inventar-Viewer (nur Ansicht)
"""

import io
import json
import os
import re
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from openpyxl.utils import get_column_letter
from pyproj import Transformer

# -----------------------------------------------
# KONSTANTEN
# -----------------------------------------------
SHEET_NAME = "Lp"
CATEGORY_ROW = 53
HEADER_ROW = 55
DATA_START_ROW = 56
DATA_START_COL = 2

ID_COL = "Lichtpunkt-Nr."
ID_COL_NEU = "Lichtpunkt-Nr. neu"
ID_COL_PROJEKT = "Lichtpunkt-Nr. Projekt"
COORD_X_COL = "Koordinate X"
COORD_Y_COL = "Koordinate Y"
COORD_COL = "GPS Koordinaten (Breite, Länge)"  # Altformat, nur noch Fallback
COORD_COLS = (COORD_X_COL, COORD_Y_COL, COORD_COL)
STREET_COL = "Strasse"

# LV95-Wertebereiche — E und N ueberlappen nicht, daher eindeutig zuordenbar
LV95_E_RANGE = (2_400_000, 2_900_000)
LV95_N_RANGE = (1_000_000, 1_400_000)

MAX_FILTER_UNIQUE = 150
EMPTY_VALS = frozenset({"", "None", "nan"})

HEX_COLORS = [
    "#e74c3c", "#3498db", "#27ae60", "#9b59b6", "#f39c12",
    "#c0392b", "#5f9ea0", "#1a5276", "#1e8449", "#e6b0aa",
    "#f5cba7", "#85c1e9", "#82e0aa", "#95a5a6", "#2c3e50",
    "#e91e90",
]

_TRANSFORMER = Transformer.from_crs("EPSG:2056", "EPSG:4326", always_xy=True)


# -----------------------------------------------
# KOORDINATEN
# -----------------------------------------------
def _to_number(series):
    """Text -> float. Akzeptiert Tausender-Apostroph/Leerzeichen und Dezimalkomma."""
    s = (
        series.astype(str).str.strip()
        .str.replace(r"[\s'’`]", "", regex=True)
        .str.replace(",", ".", regex=False)
    )
    return pd.to_numeric(s, errors="coerce")


def _in_range(s, rng):
    return s.notna() & (s > rng[0]) & (s < rng[1])


def _coords_from_xy(df):
    """LV95 aus 'Koordinate X' / 'Koordinate Y' -> (east, north) Serien.

    Ob X Ost oder Nord ist (Schweizer Vermessung: Y=Ost, X=Nord; GIS oft umgekehrt),
    wird pro Zeile am Wertebereich erkannt.
    """
    x = _to_number(df[COORD_X_COL])
    y = _to_number(df[COORD_Y_COL])
    east = x.where(_in_range(x, LV95_E_RANGE), y.where(_in_range(y, LV95_E_RANGE)))
    north = y.where(_in_range(y, LV95_N_RANGE), x.where(_in_range(x, LV95_N_RANGE)))
    return east, north


def _coords_from_legacy(df):
    """Altformat 'GPS Koordinaten': ein Zellwert mit zwei Zahlen, LV95 oder WGS84.

    Liefert (lat, lon) fuer WGS84-Zeilen und (east, north) fuer LV95-Zeilen.
    """
    raw = df[COORD_COL].str.strip()
    extracted = raw.str.extract(r"(-?\d[\d.]*)\s*[,;\s]+\s*(-?\d[\d.]*)")
    a = pd.to_numeric(extracted[0], errors="coerce")
    b = pd.to_numeric(extracted[1], errors="coerce")
    is_lv95 = _in_range(a, LV95_E_RANGE) & _in_range(b, LV95_N_RANGE)
    is_wgs = (
        a.notna() & b.notna() & ~is_lv95
        & (a >= -90) & (a <= 90) & (b >= -180) & (b <= 180)
    )
    return a.where(is_wgs), b.where(is_wgs), a.where(is_lv95), b.where(is_lv95)


def _add_coordinates(df):
    """Setzt _lat/_lon (WGS84 fuer Leaflet). Quelle primaer LV95 aus X/Y,
    fuer Zeilen ohne gueltige X/Y faellt es auf die alte GPS-Spalte zurueck."""
    nan = pd.Series(float("nan"), index=df.index)
    lat, lon, east, north = nan.copy(), nan.copy(), nan.copy(), nan.copy()

    if COORD_X_COL in df.columns and COORD_Y_COL in df.columns:
        east, north = _coords_from_xy(df)

    if COORD_COL in df.columns:
        missing = east.isna() | north.isna()
        l_lat, l_lon, l_east, l_north = _coords_from_legacy(df)
        lat, lon = l_lat.where(missing), l_lon.where(missing)
        east = east.where(~missing, l_east)
        north = north.where(~missing, l_north)

    ok = east.notna() & north.notna()
    if ok.any():
        lons, lats = _TRANSFORMER.transform(east[ok].values, north[ok].values)
        lat[ok] = lats
        lon[ok] = lons

    df["_lat"] = lat.astype(float)
    df["_lon"] = lon.astype(float)
    return df


# -----------------------------------------------
# EXCEL LESEN — XML-Struktur + Calamine-Daten
# -----------------------------------------------
_XLSX_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_RELS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_CELL_REF_RE = re.compile(r"^([A-Z]+)(\d+)$")


def _col_letter_to_index(col_str):
    """Convert column letter(s) to 1-based index. A=1, Z=26, AA=27."""
    result = 0
    for c in col_str:
        result = result * 26 + (ord(c) - 64)
    return result


def _parse_cell_ref(ref):
    """Parse 'B53' → (col_index, row_number)."""
    m = _CELL_REF_RE.match(ref)
    if m:
        return _col_letter_to_index(m.group(1)), int(m.group(2))
    return None, None


def _parse_merge_ref(ref):
    """Parse 'B53:F55' → (min_col, min_row, max_col, max_row)."""
    parts = ref.split(":")
    if len(parts) != 2:
        return None
    c1, r1 = _parse_cell_ref(parts[0])
    c2, r2 = _parse_cell_ref(parts[1])
    if c1 is not None and c2 is not None:
        return c1, r1, c2, r2
    return None


def _find_sheet_path(zf, sheet_name):
    """Find the worksheet XML path for a given sheet name."""
    # Parse workbook.xml to find sheet name → rId
    wb_xml = zf.read("xl/workbook.xml")
    wb_root = ET.fromstring(wb_xml)
    rid = None
    sheet_names = []
    for sheet_el in wb_root.iter(f"{{{_XLSX_NS}}}sheet"):
        name = sheet_el.get("name")
        sheet_names.append(name)
        if name == sheet_name:
            # r:id attribute uses a different namespace
            for attr, val in sheet_el.attrib.items():
                if attr.endswith("}id") or attr == "r:id":
                    rid = val
                    break
    if rid is None:
        raise ValueError(
            f"Arbeitsblatt '{sheet_name}' nicht gefunden. "
            f"Vorhandene Blaetter: {', '.join(sheet_names)}"
        )
    # Parse workbook.xml.rels to find rId → file path
    rels_xml = zf.read("xl/_rels/workbook.xml.rels")
    rels_root = ET.fromstring(rels_xml)
    for rel in rels_root.iter(f"{{{_RELS_NS}}}Relationship"):
        if rel.get("Id") == rid:
            target = rel.get("Target")
            # Target is relative to xl/ directory
            return f"xl/{target}" if not target.startswith("/") else target.lstrip("/")
    raise ValueError(f"Relationship '{rid}' for sheet '{sheet_name}' not found.")


def _parse_shared_strings(zf):
    """Parse xl/sharedStrings.xml into a list of strings."""
    try:
        ss_xml = zf.read("xl/sharedStrings.xml")
    except KeyError:
        return []
    strings = []
    for _, elem in ET.iterparse(io.BytesIO(ss_xml)):
        tag = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag
        if tag == "si":
            # Concatenate all <t> elements (handles both plain and rich text)
            text_parts = []
            for t_el in elem.iter(f"{{{_XLSX_NS}}}t"):
                if t_el.text:
                    text_parts.append(t_el.text)
            strings.append("".join(text_parts))
            elem.clear()
    return strings


def _get_cell_value(cell_elem, shared_strings):
    """Extract a cell's string value from its XML element."""
    cell_type = cell_elem.get("t")
    v_el = cell_elem.find(f"{{{_XLSX_NS}}}v")

    if cell_type == "s" and v_el is not None and v_el.text:
        # Shared string reference
        idx = int(v_el.text)
        return shared_strings[idx] if idx < len(shared_strings) else None
    elif cell_type == "inlineStr":
        # Inline string
        is_el = cell_elem.find(f"{{{_XLSX_NS}}}is")
        if is_el is not None:
            parts = [t.text or "" for t in is_el.iter(f"{{{_XLSX_NS}}}t")]
            return "".join(parts) or None
    elif v_el is not None and v_el.text:
        return v_el.text
    return None


def _read_structure(filepath):
    """
    Pass 1 — direct XML parsing.
    Opens the xlsx/xlsm as a ZIP and reads only the minimal XML needed
    for headers, categories, hidden cols/rows. Orders of magnitude faster
    than openpyxl.load_workbook() which parses every cell and style.
    """
    with zipfile.ZipFile(filepath, "r") as zf:
        sheet_path = _find_sheet_path(zf, SHEET_NAME)
        shared_strings = _parse_shared_strings(zf)

        hidden_cols = set()
        hidden_rows = set()
        row_cells = {}  # {row_num: {col_index: str_value}}
        merge_refs = []

        with zf.open(sheet_path) as sheet_file:
            for event, elem in ET.iterparse(sheet_file, events=("end",)):
                tag = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag

                if tag == "col":
                    if elem.get("hidden") in ("1", "true"):
                        min_c = int(elem.get("min"))
                        max_c = int(elem.get("max"))
                        hidden_cols.update(range(min_c, max_c + 1))
                    elem.clear()

                elif tag == "row":
                    r = elem.get("r")
                    if r is not None:
                        row_num = int(r)
                        if elem.get("hidden") in ("1", "true"):
                            hidden_rows.add(row_num)
                        if row_num in (CATEGORY_ROW, HEADER_ROW):
                            cells = {}
                            for c_el in elem.iter(f"{{{_XLSX_NS}}}c"):
                                ref = c_el.get("r")
                                if ref:
                                    ci, _ = _parse_cell_ref(ref)
                                    val = _get_cell_value(c_el, shared_strings)
                                    if ci is not None and val is not None:
                                        cells[ci] = val
                            row_cells[row_num] = cells
                    elem.clear()

                elif tag == "mergeCell":
                    ref = elem.get("ref")
                    if ref:
                        merge_refs.append(ref)
                    elem.clear()

    # Build category_by_col from merged cells overlapping row 53
    cat_row_cells = row_cells.get(CATEGORY_ROW, {})
    category_by_col = {}

    for ref in merge_refs:
        parsed = _parse_merge_ref(ref)
        if parsed is None:
            continue
        min_col, min_row, max_col, max_row = parsed
        if min_row <= CATEGORY_ROW <= max_row:
            # Value is in the top-left cell of the merge
            val = cat_row_cells.get(min_col)
            if val and str(val).strip():
                name = str(val).strip()
                for c in range(min_col, max_col + 1):
                    if c >= DATA_START_COL:
                        category_by_col[c] = name

    # Fill non-merged category cells
    for c, val in cat_row_cells.items():
        if c >= DATA_START_COL and c not in category_by_col:
            if val and str(val).strip():
                category_by_col[c] = str(val).strip()

    # Headers from row 55
    hdr_row_cells = row_cells.get(HEADER_ROW, {})
    headers = {}
    for c, val in hdr_row_cells.items():
        if c >= DATA_START_COL and val is not None and str(val).strip():
            headers[c] = re.sub(r"\s*\n\s*", " ", str(val).strip()).strip()

    # Remove hidden columns from headers and categories
    for ci in hidden_cols:
        headers.pop(ci, None)
        category_by_col.pop(ci, None)

    return headers, category_by_col, hidden_rows


def _read_data(filepath, headers, hidden_rows):
    """
    Pass 2 — calamine (Rust-based reader).
    Parses Excel XML in compiled code with no Python overhead per cell,
    making it 10-30x faster than openpyxl for data extraction.
    """
    from python_calamine import CalamineWorkbook

    sorted_col_items = sorted(headers.items())
    col_indices = [c for c, _ in sorted_col_items]
    header_list = [h for _, h in sorted_col_items]

    wb = CalamineWorkbook.from_path(filepath)

    if SHEET_NAME not in wb.sheet_names:
        raise ValueError(
            f"Arbeitsblatt '{SHEET_NAME}' nicht gefunden. "
            f"Vorhandene Blaetter: {', '.join(wb.sheet_names)}"
        )

    # skip_empty_area=False: Raster beginnt bei A1. Standardmaessig schneidet calamine
    # leere Randzeilen/-spalten ab, dann waeren Zeilen- und Spaltenindizes verschoben.
    all_rows = wb.get_sheet_by_name(SHEET_NAME).to_python(skip_empty_area=False)

    data_rows = []
    for row_num, row_tuple in enumerate(
        all_rows[DATA_START_ROW - 1:], start=DATA_START_ROW
    ):
        if row_num in hidden_rows:
            continue
        row_data = {"_excel_row": row_num}
        has_value = False
        for c, header in zip(col_indices, header_list):
            val = row_tuple[c - 1] if c - 1 < len(row_tuple) else None
            if val == "":
                val = None
            elif isinstance(val, float) and val.is_integer():
                # Excel speichert Ganzzahlen als float -> "30" statt "30.0".
                # Direkt als Text, sonst macht pandas gemischte Spalten wieder zu float.
                val = str(int(val))
            elif val is not None:
                val = str(val)
            row_data[header] = val
            if val is not None:
                has_value = True
        if has_value:
            data_rows.append(row_data)

    return data_rows


@st.cache_data(show_spinner="Daten werden geladen ...")
def read_excel_structure(filepath, file_mtime):
    headers, category_by_col, hidden_rows = _read_structure(filepath)
    data_rows = _read_data(filepath, headers, hidden_rows)

    df = pd.DataFrame(data_rows)
    if df.empty:
        df = pd.DataFrame(
            columns=list(headers.values()) + ["_excel_row"]
        )

    # Vectorised normalisation — single pass, no Python-level lambda per cell
    data_cols = [c for c in df.columns if not c.startswith("_")]
    df[data_cols] = df[data_cols].fillna("").astype(str)

    # Drop rows where ALL three ID columns are empty.
    # A row is kept if any one of the three carries a value.
    def _col_empty(col):
        if col in df.columns:
            return df[col].str.strip().isin(EMPTY_VALS)
        return pd.Series(True, index=df.index)

    drop_mask = _col_empty(ID_COL) & _col_empty(ID_COL_NEU) & _col_empty(ID_COL_PROJEKT)
    df = df[~drop_mask].reset_index(drop=True)

    # Category and col_info mapping
    sorted_col_keys = sorted(headers.keys())
    cat_map = {}
    ordered_categories = []
    seen_cats: set = set()
    for ci in sorted_col_keys:
        h = headers[ci]
        cat = category_by_col.get(ci, "Sonstige")
        cat_map[h] = cat
        if cat not in seen_cats:
            ordered_categories.append(cat)
            seen_cats.add(cat)

    col_info = {
        headers[ci]: {
            "col_index": ci,
            "category": category_by_col.get(ci, "Sonstige"),
        }
        for ci in sorted_col_keys
    }

    # Batch coordinate conversion
    df = _add_coordinates(df)

    return df, cat_map, ordered_categories, col_info


# -----------------------------------------------
# DATEI-BROWSER
# -----------------------------------------------
def render_file_selector():
    st.title("Leuchten GIS Dashboard")
    st.markdown("### Datei auswaehlen")
    st.caption("Waehle eine `.xlsx` oder `.xlsm` Datei – auch von Netzlaufwerken (R:\\, A:\\).")
    st.markdown("---")

    available_drives = [d for d in ["R:\\", "A:\\", "C:\\", "D:\\", "E:\\"] if os.path.exists(d)]

    if available_drives:
        st.markdown("**Schnellzugriff:**")
        drive_cols = st.columns(min(len(available_drives), 5))
        for i, drive in enumerate(available_drives):
            with drive_cols[i]:
                if st.button(drive, key=f"drv_{drive}", use_container_width=True):
                    st.session_state["browser_path"] = drive
                    st.rerun()

    st.markdown("---")

    default_path = st.session_state.get(
        "browser_path", available_drives[0] if available_drives else "C:\\"
    )
    path_input = st.text_input(
        "Pfad eingeben oder einfuegen",
        value=default_path,
        placeholder="z.B. R:\\Projekte\\leuchten.xlsx",
        key="path_input_field",
    )

    if path_input and os.path.isfile(path_input):
        ext = os.path.splitext(path_input)[1].lower()
        if ext in (".xlsx", ".xlsm"):
            st.success(f"**{os.path.basename(path_input)}** gefunden")
            if st.button("Diese Datei oeffnen", type="primary", use_container_width=True):
                return path_input
            return None

    browse_dir = path_input if os.path.isdir(path_input) else os.path.dirname(path_input)
    if not browse_dir or not os.path.isdir(browse_dir):
        browse_dir = available_drives[0] if available_drives else "C:\\"

    st.session_state["browser_path"] = browse_dir

    try:
        entries = sorted(os.listdir(browse_dir), key=str.lower)
    except PermissionError:
        st.error("Zugriff verweigert.")
        return None
    except OSError as e:
        st.error(f"Fehler: {e}")
        return None

    folders = [
        e for e in entries
        if os.path.isdir(os.path.join(browse_dir, e)) and not e.startswith(".")
    ]
    files = [
        e for e in entries
        if os.path.isfile(os.path.join(browse_dir, e))
        and os.path.splitext(e)[1].lower() in (".xlsx", ".xlsm")
    ]

    parent = os.path.dirname(browse_dir.rstrip("\\/"))
    if parent and parent != browse_dir:
        if st.button(".. (uebergeordneter Ordner)", key="nav_parent"):
            st.session_state["browser_path"] = parent
            st.rerun()

    if folders:
        st.markdown(f"**Ordner** in `{browse_dir}`")
        cols_per_row = 3
        for row_start in range(0, len(folders), cols_per_row):
            cols = st.columns(cols_per_row)
            for j, folder in enumerate(folders[row_start: row_start + cols_per_row]):
                with cols[j]:
                    if st.button(folder, key=f"fld_{row_start}_{j}", use_container_width=True):
                        st.session_state["browser_path"] = os.path.join(browse_dir, folder)
                        st.rerun()

    if files:
        st.markdown(f"**Excel-Dateien** in `{browse_dir}`")
        for f in files:
            full_path = os.path.join(browse_dir, f)
            try:
                size_mb = os.path.getsize(full_path) / (1024 * 1024)
            except OSError:
                size_mb = 0
            icon = "xlsm" if f.lower().endswith(".xlsm") else "xlsx"
            col_name, col_btn = st.columns([4, 1])
            with col_name:
                st.markdown(f"**{f}**  ({size_mb:.1f} MB, {icon})")
            with col_btn:
                if st.button("Oeffnen", key=f"open_{f}", use_container_width=True):
                    return full_path

    if not folders and not files:
        st.info("Keine Ordner oder Excel-Dateien in diesem Verzeichnis.")

    return None


# -----------------------------------------------
# FARBPALETTEN
# -----------------------------------------------
def build_color_map(series):
    unique_vals = sorted(series.dropna().unique().tolist(), key=str)
    return {str(v): HEX_COLORS[i % len(HEX_COLORS)] for i, v in enumerate(unique_vals)}


# -----------------------------------------------
# KARTE (Leaflet)
# -----------------------------------------------
@st.cache_data(show_spinner="Karte wird erstellt ...")
def build_map_html(
    markers_json, color_col, color_map_json, center_lat, center_lon, zoom,
):
    # Alle Parameter sind Teil des Cache-Schluessels (markers_json enthaelt die gefilterten Punkte)
    html = f"""<!DOCTYPE html>
<html><head>
    <meta charset="utf-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <link rel="stylesheet" href="https://unpkg.com/leaflet@1.9.4/dist/leaflet.css"/>
    <style>
        :root {{
            --bg: #ffffff; --fg: #222222; --fg-muted: #666666;
            --border: #cccccc; --border-light: #dddddd;
            --shadow: rgba(0,0,0,.15);
            --legend-bg: white; --legend-title: #333333;
            --popup-bg: #ffffff; --popup-fg: #222222;
            --tooltip-bg: #ffffff; --tooltip-fg: #222222;
        }}
        @media (prefers-color-scheme: dark) {{
            :root {{
                --bg: #1a1a2e; --fg: #e0e0e0; --fg-muted: #a0a0b0;
                --border: #444466; --border-light: #3a3a55;
                --shadow: rgba(0,0,0,.4);
                --legend-bg: #222244; --legend-title: #d0d0e0;
                --popup-bg: #222244; --popup-fg: #e0e0e0;
                --tooltip-bg: #2a2a4a; --tooltip-fg: #e0e0e0;
            }}
        }}
        html, body {{ margin: 0; padding: 0; height: 100%; }}
        #map {{ width: 100%; height: 100%; }}
        .legend {{
            position: absolute; bottom: 20px; right: 10px; z-index: 1000;
            background: var(--legend-bg); padding: 12px 16px; border-radius: 8px;
            border: 1px solid var(--border); font-family: sans-serif; font-size: 12px;
            max-height: 280px; overflow-y: auto;
            box-shadow: 3px 3px 8px var(--shadow); color: var(--fg);
        }}
        .legend b {{ font-size: 12px; color: var(--legend-title); }}
        .legend-item {{ margin: 3px 0; display: flex; align-items: center; }}
        .legend-picker {{
            width: 24px; height: 24px; border: 1px solid var(--border);
            border-radius: 50%; cursor: pointer; padding: 0;
            margin-right: 8px; flex-shrink: 0;
            -webkit-appearance: none; appearance: none;
            background: none; overflow: hidden;
        }}
        .legend-picker::-webkit-color-swatch-wrapper {{ padding: 0; }}
        .legend-picker::-webkit-color-swatch {{ border: none; border-radius: 50%; }}
        .legend-picker::-moz-color-swatch {{ border: none; border-radius: 50%; }}
        .popup-content {{
            font-family: sans-serif; font-size: 12px;
            max-height: 320px; overflow-y: auto;
            min-width: 240px; padding-right: 6px;
        }}
        .popup-content b.title {{ font-size: 14px; }}
        .popup-content .street {{ color: var(--fg-muted); }}
        .popup-content hr {{ margin: 5px 0; border: none; border-top: 1px solid var(--border-light); }}
        /* Leaflet popup/tooltip dark mode overrides */
        .leaflet-popup-content-wrapper {{
            background: var(--popup-bg) !important;
            color: var(--popup-fg) !important;
        }}
        .leaflet-popup-tip {{ background: var(--popup-bg) !important; }}
        .leaflet-tooltip {{
            background: var(--tooltip-bg) !important;
            color: var(--tooltip-fg) !important;
            border-color: var(--border) !important;
        }}
        .leaflet-tooltip-top:before {{ border-top-color: var(--tooltip-bg) !important; }}
        .leaflet-tooltip-bottom:before {{ border-bottom-color: var(--tooltip-bg) !important; }}
        .leaflet-tooltip-left:before {{ border-left-color: var(--tooltip-bg) !important; }}
        .leaflet-tooltip-right:before {{ border-right-color: var(--tooltip-bg) !important; }}
    </style>
</head><body>
    <div id="map"></div>
    <script src="https://unpkg.com/leaflet@1.9.4/dist/leaflet.js"></script>
    <script>
    (function() {{
        const map = L.map('map', {{
            center: [{center_lat}, {center_lon}],
            zoom: {zoom},
            preferCanvas: true
        }});
        L.tileLayer('https://{{s}}.tile.openstreetmap.org/{{z}}/{{x}}/{{y}}.png', {{
            attribution: '&copy; OpenStreetMap', maxZoom: 19,
        }}).addTo(map);

        const markers = {markers_json};
        const colorMap = {color_map_json};
        const colorCol = {json.dumps(color_col)};

        function escHtml(s) {{
            return String(s)
                .replace(/&/g, '&amp;')
                .replace(/</g, '&lt;')
                .replace(/>/g, '&gt;')
                .replace(/"/g, '&quot;');
        }}

        const group = L.featureGroup();
        const markersByValue = {{}};

        for (let i = 0; i < markers.length; i++) {{
            const m = markers[i];
            const marker = L.circleMarker([m.lat, m.lon], {{
                radius: 6, color: m.color, fillColor: m.color,
                fillOpacity: 0.85, weight: 1,
            }});
            if (m.catVal !== undefined && m.catVal !== '') {{
                if (!markersByValue[m.catVal]) markersByValue[m.catVal] = [];
                markersByValue[m.catVal].push(marker);
            }}
            marker.bindTooltip(
                escHtml(m.id) + (m.street ? ' | ' + escHtml(m.street) : ''),
                {{ direction: 'top', offset: [0, -8] }}
            );
            marker.on('click', function() {{
                if (!this._popupBuilt) {{
                    let html = '<div class="popup-content">';
                    html += '<b class="title">' + escHtml(m.id) + '</b>';
                    html += '<br><span class="street">' + escHtml(m.street || '') + '</span>';
                    html += '<hr>';
                    const fields = m.fields;
                    for (const key in fields) {{
                        const val = fields[key];
                        if (val && val !== 'nan' && val.trim() !== '') {{
                            html += '<b>' + escHtml(key) + ':</b> ' + escHtml(val) + '<br>';
                        }}
                    }}
                    html += '</div>';
                    this.bindPopup(html, {{ maxWidth: 320 }});
                    this._popupBuilt = true;
                }}
                this.openPopup();
            }});
            marker.addTo(group);
        }}

        group.addTo(map);
        if (markers.length > 0) {{ map.fitBounds(group.getBounds().pad(0.05)); }}

        if (colorCol && Object.keys(colorMap).length > 0) {{
            const legend = L.control({{ position: 'bottomright' }});
            legend.onAdd = function() {{
                const div = L.DomUtil.create('div', 'legend');
                div.innerHTML = '<b>' + escHtml(colorCol) + '</b><br><br>';
                for (const [val, col] of Object.entries(colorMap)) {{
                    const row = document.createElement('div');
                    row.className = 'legend-item';
                    const picker = document.createElement('input');
                    picker.type = 'color'; picker.className = 'legend-picker';
                    picker.value = col;
                    picker.title = 'Farbe aendern: ' + escHtml(val);
                    picker.addEventListener('input', function() {{
                        const newColor = this.value;
                        const targets = markersByValue[val] || [];
                        for (let j = 0; j < targets.length; j++) {{
                            targets[j].setStyle({{ color: newColor, fillColor: newColor }});
                        }}
                    }});
                    const label = document.createElement('span');
                    label.textContent = val;
                    row.appendChild(picker);
                    row.appendChild(label);
                    div.appendChild(row);
                }}
                L.DomEvent.disableClickPropagation(div);
                L.DomEvent.disableScrollPropagation(div);
                return div;
            }};
            legend.addTo(map);
        }}
    }})();
    </script>
</body></html>"""
    return html


def prepare_map_data(df, color_col):
    valid = df.dropna(subset=["_lat", "_lon"]).copy()
    if valid.empty:
        return None

    lats = tuple(valid["_lat"].astype(float).tolist())
    lons = tuple(valid["_lon"].astype(float).tolist())
    # Titel im Popup: erste vorhandene der drei ID-Spalten
    id_series = pd.Series("", index=valid.index)
    for c in (ID_COL_PROJEKT, ID_COL_NEU, ID_COL):
        if c in valid.columns:
            val = valid[c].str.strip()
            id_series = val.where(~val.isin(EMPTY_VALS), id_series)
    ids = tuple(id_series.replace("", "-").tolist())
    streets = (
        tuple(valid[STREET_COL].fillna("").tolist())
        if STREET_COL in valid.columns
        else ("",) * len(valid)
    )

    color_map = {}
    cat_vals = ("",) * len(valid)
    if color_col and color_col in valid.columns:
        color_map = build_color_map(valid[color_col])
        cat_vals = tuple(valid[color_col].fillna("").astype(str).tolist())
        colors = tuple(
            valid[color_col].fillna("").astype(str)
            .map(lambda v, cm=color_map: cm.get(v, "#808080"))
            .tolist()
        )
    else:
        colors = ("#808080",) * len(valid)

    skip_cols = {"_lat", "_lon", "_excel_row", COORD_COL}
    display_cols = [c for c in df.columns if c not in skip_cols and not c.startswith("_")]

    # to_dict('records') is 3-5x faster than iterrows() for building popup dicts
    records = valid[display_cols].to_dict("records")
    popup_rows = tuple(
        {
            "Excel-Zeile": str(int(excel_row)),
            **{
                k: str(v)
                for k, v in rec.items()
                if pd.notna(v) and str(v).strip() not in EMPTY_VALS
            },
        }
        for excel_row, rec in zip(valid["_excel_row"], records)
    )

    markers_json = json.dumps(
        [
            {
                "lat": lat, "lon": lon, "id": mid, "street": street,
                "color": color, "catVal": cat_val, "fields": popup,
            }
            for lat, lon, mid, street, color, cat_val, popup
            in zip(lats, lons, ids, streets, colors, cat_vals, popup_rows)
        ],
        ensure_ascii=False,
    )

    center_lat = sum(lats) / len(lats)
    center_lon = sum(lons) / len(lons)

    return {
        "markers_json": markers_json,
        "color_col": color_col,
        "color_map_json": json.dumps(color_map, ensure_ascii=False),
        "center_lat": center_lat, "center_lon": center_lon, "zoom": 14,
    }


# -----------------------------------------------
# HILFSFUNKTIONEN
# -----------------------------------------------
def get_columns_by_category(cat_map, ordered_categories):
    result = {cat: [] for cat in ordered_categories}
    for header, cat in cat_map.items():
        if cat in result:
            result[cat].append(header)
    return result


def is_filterable(df, col):
    if col not in df.columns:
        return False
    non_empty = df[col][~df[col].str.strip().isin(EMPTY_VALS)]
    return 1 < non_empty.nunique() <= MAX_FILTER_UNIQUE


# -----------------------------------------------
# STREAMLIT APP
# -----------------------------------------------
st.set_page_config(
    page_title="Leuchten GIS",
    page_icon="💡",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown("""
<style>
    /* --- Light mode sidebar --- */
    [data-testid="stSidebar"] { background: #f0f2f8; }
    [data-testid="stSidebar"] .stSelectbox label,
    [data-testid="stSidebar"] .stMultiSelect label,
    [data-testid="stSidebar"] .stTextInput label { color: #555e70 !important; font-size: 11px !important; }
    div[data-testid="metric-container"] {
        background: #f5f7ff;
        border-left: 4px solid #3a5bff;
        border-radius: 6px;
        padding: 8px 12px;
    }
    /* --- Dark mode overrides --- */
    @media (prefers-color-scheme: dark) {
        [data-testid="stSidebar"] { background: #1a1f2e; }
        [data-testid="stSidebar"] * { color: #e8eaf0 !important; }
        [data-testid="stSidebar"] .stSelectbox label,
        [data-testid="stSidebar"] .stMultiSelect label,
        [data-testid="stSidebar"] .stTextInput label { color: #8892b0 !important; font-size: 11px !important; }
        div[data-testid="metric-container"] {
            background: #1e2433;
            border-left-color: #5a7bff;
        }
    }
    [data-testid="stTabs"] iframe {
        height: calc(100vh - 280px) !important;
        min-height: 500px;
    }
</style>
""", unsafe_allow_html=True)


# --- DATEIAUSWAHL ---
if "filepath" not in st.session_state or st.session_state.filepath is None:
    selected = render_file_selector()
    if selected:
        st.session_state.filepath = selected
        st.rerun()
    st.stop()

# --- DATEN LADEN ---
filepath = st.session_state.filepath

_LOAD_KEYS = ["filepath", "df", "cat_map", "ordered_categories", "col_info", "filterable_cols"]

if "df" not in st.session_state:
    try:
        df, cat_map, ordered_categories, col_info = read_excel_structure(
            filepath, os.path.getmtime(filepath)
        )
        id_cols_found = [c for c in (ID_COL, ID_COL_NEU, ID_COL_PROJEKT) if c in col_info]
        if not id_cols_found:
            raise ValueError(
                f"Keine ID-Spalte in Zeile {HEADER_ROW} gefunden "
                f"(erwartet: '{ID_COL}', '{ID_COL_NEU}' oder '{ID_COL_PROJEKT}'). "
                f"Gefundene Spalten: {', '.join(col_info) or 'keine'}"
            )
        st.session_state.df = df
        st.session_state.cat_map = cat_map
        st.session_state.ordered_categories = ordered_categories
        st.session_state.col_info = col_info
        # Compute filterable columns once per file load, not on every rerun
        _all_init = [c for c in df.columns if not c.startswith("_")]
        _always_filterable = {ID_COL, ID_COL_NEU, ID_COL_PROJEKT}
        st.session_state.filterable_cols = [
            c for c in _all_init
            if is_filterable(df, c) or c in _always_filterable
        ]
    except Exception as e:
        st.error(f"Fehler beim Laden: {e}")
        if st.button("Andere Datei waehlen"):
            for key in _LOAD_KEYS:
                st.session_state.pop(key, None)
            st.rerun()
        st.stop()

df = st.session_state.df
cat_map = st.session_state.cat_map
ordered_categories = st.session_state.ordered_categories
col_info = st.session_state.col_info
filterable_cols = st.session_state.get("filterable_cols", [])

# Computed once per rerun from stable session state — not reloaded from Excel
cols_by_cat = get_columns_by_category(cat_map, ordered_categories)
all_data_cols = [c for c in df.columns if not c.startswith("_")]


# --- SIDEBAR ---
with st.sidebar:
    st.markdown("## Leuchten GIS")
    st.caption(os.path.basename(filepath))
    st.markdown("---")

    if st.button("Andere Datei oeffnen", use_container_width=True):
        for key in _LOAD_KEYS:
            st.session_state.pop(key, None)
        st.cache_data.clear()
        st.rerun()

    if st.button("Daten neu laden", use_container_width=True):
        for key in _LOAD_KEYS[1:]:  # keep filepath
            st.session_state.pop(key, None)
        st.cache_data.clear()
        st.rerun()

    st.markdown("---")

    color_col = st.selectbox(
        "Farbcodierung nach",
        options=filterable_cols if filterable_cols else all_data_cols,
    )

    st.markdown("---")
    st.markdown("### Filter")

    # Accumulate boolean mask — no df.copy()
    mask = pd.Series(True, index=df.index)

    for cat in ordered_categories:
        cat_cols = cols_by_cat.get(cat, [])
        filter_cols_in_cat = [c for c in cat_cols if c in filterable_cols]
        if not filter_cols_in_cat:
            continue

        with st.expander(cat, expanded=False):
            for col in filter_cols_in_cat:
                unique_vals = sorted(
                    df[col][~df[col].str.strip().isin(EMPTY_VALS)].unique().tolist(),
                    key=str,
                )
                if not unique_vals:
                    continue
                sel = st.multiselect(col, unique_vals, key=f"filter_{col}")
                if sel:
                    mask &= df[col].isin(sel)

    st.markdown("---")

    search = st.text_input("Suche (Strasse / Lichtpunkt-Nr.)", "")
    if search:
        s = search.strip().lower()
        search_cols = [c for c in (STREET_COL, ID_COL, ID_COL_NEU, ID_COL_PROJEKT) if c in df.columns]
        hit = pd.Series(False, index=df.index)
        for c in search_cols:
            # regex=False: Eingaben wie "(" oder "[" sind normaler Text
            hit |= df[c].str.lower().str.contains(s, regex=False, na=False)
        mask &= hit

    filtered = df[mask]

    st.markdown("---")
    valid_count = int(filtered["_lat"].notna().sum())
    st.caption(f"**{len(filtered)}** Leuchten | **{valid_count}** mit Koordinaten")

    def _export_bytes(data=filtered):
        buf = io.BytesIO()
        data[[c for c in data.columns if not c.startswith("_")]].to_excel(
            buf, index=False, sheet_name=SHEET_NAME
        )
        return buf.getvalue()

    _src_name = os.path.splitext(os.path.basename(filepath))[0]
    st.download_button(
        "Export (gefiltert)",
        data=_export_bytes,
        file_name=f"{_src_name}_export_{datetime.now():%Y%m%d_%H%M}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        on_click="ignore",
        use_container_width=True,
        help="Laedt die gefilterten Leuchten als Excel-Datei herunter (Download-Ordner).",
    )


# --- HAUPTBEREICH ---
st.title("Leuchten GIS Dashboard")
st.caption(os.path.basename(filepath))

m1, m2, m3 = st.columns(3)
m1.metric("Leuchten gesamt", len(df))
m2.metric("Gefiltert", len(filtered))
m3.metric("Mit Koordinaten", valid_count)

st.caption(
    "Nur Ansicht – Aenderungen bitte direkt in Excel vornehmen und danach "
    "in der Seitenleiste «Daten neu laden» klicken."
)

st.markdown("---")

tab_map, tab_table, tab_debug = st.tabs(["Karte", "Tabelle", "Diagnose"])

# -- TAB: KARTE --
with tab_map:
    if filtered["_lat"].notna().sum() == 0:
        st.warning("Keine Leuchten mit gueltigen Koordinaten in der aktuellen Auswahl.")
    else:
        map_data = prepare_map_data(filtered, color_col)
        if map_data:
            map_html = build_map_html(**map_data)
            components.html(map_html, height=860, scrolling=False)

# -- TAB: TABELLE --
with tab_table:
    display_cols = [c for c in filtered.columns if not c.startswith("_")]
    st.dataframe(
        filtered[["_excel_row"] + display_cols].rename(columns={"_excel_row": "Excel-Zeile"}),
        use_container_width=True,
        height=520,
        hide_index=True,
    )

# -- TAB: DIAGNOSE --
with tab_debug:
    st.markdown("#### Diagnose")

    st.markdown(f"**Datei:** `{filepath}`")
    st.markdown(f"**Arbeitsblatt:** `{SHEET_NAME}`")
    st.markdown(f"**Kategorien:** {len(ordered_categories)}")
    st.markdown(f"**Parameter (Spalten):** {len(col_info)}")
    st.markdown(f"**Datensaetze:** {len(df)}")

    st.markdown("---")
    st.markdown("**Kategorien und zugehoerige Spalten:**")
    for cat in ordered_categories:
        cat_cols = cols_by_cat.get(cat, [])
        with st.expander(f"{cat} ({len(cat_cols)} Spalten)", expanded=False):
            for c in cat_cols:
                col_letter = get_column_letter(col_info[c]["col_index"])
                st.write(f"  {col_letter}: `{c}`")

    st.markdown("---")
    for check_col, label in [
        (COORD_X_COL, "Koordinate X (LV95)"),
        (COORD_Y_COL, "Koordinate Y (LV95)"),
        (COORD_COL, "GPS-Spalte (Altformat, Fallback)"),
        (ID_COL, "ID-Spalte"),
        (STREET_COL, "Strasse-Spalte"),
    ]:
        if check_col in col_info:
            letter = get_column_letter(col_info[check_col]["col_index"])
            status = f"gefunden in Spalte {letter}"
        else:
            status = "NICHT gefunden"
        st.markdown(f"**{label}** `{check_col}`: {status}")

    present_coord_cols = [c for c in COORD_COLS if c in df.columns]
    if present_coord_cols:
        st.markdown("---")
        has_raw = pd.Series(False, index=df.index)
        for c in present_coord_cols:
            has_raw |= ~df[c].str.strip().isin(EMPTY_VALS)
        invalid = df[has_raw & df["_lat"].isna()]

        st.markdown("**Koordinaten roh und umgerechnet (erste 3):**")
        st.dataframe(
            df.loc[df["_lat"].notna(), present_coord_cols + ["_lat", "_lon"]].head(3),
            use_container_width=True,
        )
        st.markdown(f"**Zeilen mit Koordinaten-Eintrag, aber nicht interpretierbar:** {len(invalid)}")
        if len(invalid):
            id_cols = [c for c in (ID_COL, STREET_COL) if c in df.columns]
            st.dataframe(
                invalid[["_excel_row"] + id_cols + present_coord_cols]
                .rename(columns={"_excel_row": "Excel-Zeile"}),
                use_container_width=True,
                hide_index=True,
            )

