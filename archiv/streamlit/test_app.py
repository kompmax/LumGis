"""
Regressionstests fuer app.py — ohne pytest lauffaehig:

    python archiv/streamlit/test_app.py

Erzeugt Testdateien in einem Temp-Ordner (Excel-Dateien sind per .gitignore
ausgeschlossen) und faehrt die App mit Streamlits AppTest durch.
"""

import logging
import os
import sys
import tempfile
import warnings

import openpyxl
from streamlit.testing.v1 import AppTest

warnings.filterwarnings("ignore")
logging.disable(logging.WARNING)

APP = os.path.join(os.path.dirname(os.path.abspath(__file__)), "app.py")

HEADERS = {
    2: "Lichtpunkt-Nr.", 3: "Lichtpunkt-Nr. neu", 4: "Lichtpunkt-Nr. Projekt",
    5: "Strasse", 6: "Intern versteckt", 7: "Koordinate\n X", 8: "Bemerkung",
    9: "Koordinate\n Y", 10: "Leistung W", 11: "Mastnummer",
}
CATEGORIES = [("Identifikation", 2, 4), ("Standort", 5, 9), ("Leuchte", 10, 11)]

# (ID, ID neu, Strasse, X, Y, Leistung, Mastnummer)
ROWS = [
    ("LP-0001", None, "Bahnhofstrasse", 2683000.5, 1247000.5, 30, "00001"),
    ("LP-0002", None, "Dorfstrasse (alt)", 1247100, 2683100, 70, "00002"),   # X/Y vertauscht
    ("LP-0003", None, "Seestrasse", "2'683'200", "1'247'200", 45, "00003"),  # Text mit Apostroph
    ("LP-0004", None, "Seestrasse", "2683300,5", "1247300,5", 45.5, "00004"),  # Dezimalkomma
    ("LP-0005", None, "Industriestr. 5", None, None, 125, "00005"),          # keine Koordinaten
    ("LP-0006", None, "Industriestr. 5", 2683400, "folgt", 30, "00006"),     # ungueltig
    (None, "N-7", "Rämistrasse", 2683500, 1247500, 30, "00007"),             # nur ID neu
    ("LP-HIDDEN", None, "Versteckt", 2683600, 1247600, 30, "00008"),         # versteckte Zeile
]


def make_file(path, layout):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lp"
    wb.create_sheet("Andere")
    if layout == "a15":
        ws["A15"] = "Info"  # wie in den echten Dateien: A1 leer, Inhalt ab A15
    elif layout == "a1":
        ws["A1"] = "Titel"
    # layout == "leer": nichts vor Zeile 53, Spalte A komplett leer
    for name, c1, c2 in CATEGORIES:
        ws.cell(53, c1, name)
        ws.merge_cells(start_row=53, start_column=c1, end_row=53, end_column=c2)
    for c, h in HEADERS.items():
        ws.cell(55, c, h)
    ws.column_dimensions["F"].hidden = True
    for i, (lp, lp_neu, street, x, y, watt, mast) in enumerate(ROWS):
        r = 56 + i
        for c, v in zip((2, 3, 5, 6, 7, 9, 10, 11), (lp, lp_neu, street, "geheim", x, y, watt, mast)):
            if v is not None:
                ws.cell(r, c, v)
        if lp == "LP-HIDDEN":
            ws.row_dimensions[r].hidden = True
    wb.save(path)


def run_app(path, search=None):
    at = AppTest.from_file(APP, default_timeout=60)
    at.session_state["filepath"] = path
    at.run()
    if search is not None:
        at.sidebar.text_input[0].set_value(search).run()
    return at


FAILS = []


def check(name, cond, detail=""):
    print(f"  {'OK  ' if cond else 'FAIL'} {name}{'' if cond else ' -> ' + str(detail)}")
    if not cond:
        FAILS.append(name)


def test_layout(tmp, layout):
    print(f"\nLayout '{layout}':")
    path = os.path.join(tmp, f"lp_{layout}.xlsx")
    make_file(path, layout)
    at = run_app(path)
    check("keine Exception", not at.exception, [e.message for e in at.exception])
    df = at.session_state["df"]
    by_id = df.set_index(df["Lichtpunkt-Nr."].where(df["Lichtpunkt-Nr."] != "", df["Lichtpunkt-Nr. neu"]))

    check("7 sichtbare Leuchten (versteckte Zeile ignoriert)", len(df) == 7, len(df))
    check("versteckte Spalte ignoriert", "Intern versteckt" not in df.columns)
    check("Spalten nicht verschoben", by_id.loc["LP-0001", "Strasse"] == "Bahnhofstrasse",
          by_id.loc["LP-0001"].to_dict())
    check("Excel-Zeile korrekt", int(by_id.loc["LP-0001", "_excel_row"]) == 56)
    check("Ganzzahl ohne .0", by_id.loc["LP-0001", "Leistung W"] == "30", by_id.loc["LP-0001", "Leistung W"])
    check("Dezimalzahl bleibt", by_id.loc["LP-0004", "Leistung W"] == "45.5")
    check("fuehrende Nullen bleiben", by_id.loc["LP-0001", "Mastnummer"] == "00001")
    check("Zeile nur mit ID neu bleibt", "N-7" in by_id.index)
    has = by_id["_lat"].notna()
    check("Koordinaten: 5 gueltig", int(has.sum()) == 5, int(has.sum()))
    check("X/Y vertauscht erkannt", bool(has["LP-0002"]))
    check("Apostroph + Dezimalkomma", bool(has["LP-0003"] and has["LP-0004"]))
    check("ungueltig/fehlend ohne Punkt", not has["LP-0005"] and not has["LP-0006"])
    lat, lon = by_id.loc["LP-0001", ["_lat", "_lon"]]
    check("Umrechnung plausibel (Zuerich)", 47.3 < lat < 47.4 and 8.5 < lon < 8.6, (lat, lon))


def test_search(tmp):
    print("\nSuche:")
    path = os.path.join(tmp, "lp_a15.xlsx")
    for term, expected in [("(", 1), ("Dorfstrasse (alt)", 1), ("[", 0), ("industriestr. 5", 2),
                           ("N-7", 1), ("lp-000", 6)]:
        at = run_app(path, search=term)
        ok = not at.exception
        gefiltert = int(at.metric[1].value) if ok else None
        check(f"Suche {term!r} -> {expected}", ok and gefiltert == expected,
              [e.message for e in at.exception] or gefiltert)


def test_missing_ids(tmp):
    print("\nFehlende ID-Spalten:")
    path = os.path.join(tmp, "lp_ohne_id.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lp"
    ws["B55"], ws["C55"] = "Strasse", "Koordinate X"
    ws["B56"], ws["C56"] = "Weg", 2683000
    wb.save(path)
    at = run_app(path)
    msgs = " ".join(e.value for e in at.error)
    check("verstaendliche Fehlermeldung", "Keine ID-Spalte" in msgs, msgs or "keine Meldung")


def test_no_edit_tab(tmp):
    print("\nNur Ansicht:")
    at = run_app(os.path.join(tmp, "lp_a15.xlsx"))
    labels = [t.label for t in at.tabs]
    check("kein Bearbeiten-Tab", "Bearbeiten" not in labels, labels)


if __name__ == "__main__":
    with tempfile.TemporaryDirectory() as tmp:
        for layout in ("a15", "a1", "leer"):
            test_layout(tmp, layout)
        test_search(tmp)
        test_missing_ids(tmp)
        test_no_edit_tab(tmp)
    print(f"\n{'ALLE TESTS OK' if not FAILS else f'{len(FAILS)} FEHLGESCHLAGEN: ' + ', '.join(FAILS)}")
    sys.exit(1 if FAILS else 0)
